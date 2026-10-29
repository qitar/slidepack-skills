#!/usr/bin/env python3
# /// script
# requires-python = ">=3.8"
# dependencies = ["python-pptx>=0.6.21"]
# ///
"""
SlidePack template.pptx prober.

Opens a template PPTX, walks every shape on every slide (recursing into groups),
and prints the keys the template expects, for BOTH of SlidePack's binding modes:

  1. {key} placeholders — text boxes, image/video shapes, AND chart titles /
     axis titles (these are ordinary text placeholders, not chart-object fields).
  2. alt text          — tables and charts bind by the shape's alt text (descr).

It also reports each chart's type, each table's dimensions, and — with --data —
a COVERAGE check against a data.json: which template keys your data leaves
unfilled (a `{key}` would leak into the output) and which data keys the template
ignores.

Usage:
    python probe_template.py [template.pptx] [--data data.json] [--slide N[,N...]] [--all] [--json]

    python probe_template.py                       # probe ./template.pptx
    python probe_template.py deck.pptx             # probe a named file
    python probe_template.py --data data.json      # probe + coverage check
    python probe_template.py --slide 15            # only slide 15 (the `template` value)
    python probe_template.py --json                # machine-readable output

With --data, only the slides the data refers to are listed (add --all for every slide).

Exit codes:
    0  ok (probe succeeded; with --data, full coverage — nothing unfilled)
    1  coverage gap (with --data: unfilled placeholder keys remain)
    2  usage / environment error (file missing, python-pptx not installed)

Requires: python-pptx. Start it through scripts/run.py, which supplies it without
touching your own Python: python scripts/run.py probe_template.py template.pptx
"""

import argparse
import json
import os
import re
import sys

PLACEHOLDER_RE = re.compile(r"\{([^{}]+)\}")


# ----------------------------------------------------------------------------
# Encoding safety
# ----------------------------------------------------------------------------
def _make_output_encoding_safe():
    """Stop a legacy console codepage from turning a report into a traceback.

    This script prints ✓ / ✗ / ── / ▪. On Windows the console encoding is often
    a legacy codepage (cp932 on a Japanese system, cp1252 elsewhere) that cannot
    represent them, and print() raises UnicodeEncodeError mid-report. Ask for
    UTF-8; if the stream refuses, at least make it substitute rather than raise.
    """
    for stream in (sys.stdout, sys.stderr):
        try:
            stream.reconfigure(encoding="utf-8")
        except Exception:
            try:
                stream.reconfigure(errors="replace")
            except Exception:
                pass


_make_output_encoding_safe()


# ----------------------------------------------------------------------------
# Output helpers
# ----------------------------------------------------------------------------
class C:
    RED = "\033[31m"
    GREEN = "\033[32m"
    YELLOW = "\033[33m"
    CYAN = "\033[36m"
    DIM = "\033[2m"
    BOLD = "\033[1m"
    OFF = "\033[0m"


_COLOR = sys.stdout.isatty() and os.environ.get("NO_COLOR") is None


def paint(text, color):
    return f"{color}{text}{C.OFF}" if _COLOR else text


def die(msg, code=2):
    print(paint("ERROR: ", C.RED) + msg, file=sys.stderr)
    sys.exit(code)


# ----------------------------------------------------------------------------
# python-pptx access helpers
# ----------------------------------------------------------------------------
def get_alt_text(shape):
    """Return (title, descr) from the shape's cNvPr (alt text), or ('', '')."""
    try:
        from pptx.oxml.ns import qn
        cnvpr = next(shape._element.iter(qn("p:cNvPr")), None)
        if cnvpr is not None:
            return cnvpr.get("title", "") or "", cnvpr.get("descr", "") or ""
    except Exception:
        pass
    return "", ""


def placeholders_in(text):
    return PLACEHOLDER_RE.findall(text or "")


def chart_info(shape):
    """Extract chart type plus any {key} placeholders in title / axis titles."""
    info = {"chart_type": "", "title_keys": [], "axis_title_keys": [], "series": 0}
    try:
        chart = shape.chart
    except Exception:
        return info
    try:
        info["chart_type"] = getattr(chart.chart_type, "name", str(chart.chart_type))
    except Exception:
        pass
    # chart title text may contain a {key} placeholder
    try:
        if chart.has_title:
            info["title_keys"] = placeholders_in(chart.chart_title.text_frame.text)
    except Exception:
        pass
    # axis titles (category/value for bar/line/area; scatter has no category_axis)
    for axis_attr in ("category_axis", "value_axis"):
        try:
            axis = getattr(chart, axis_attr)
            if axis is not None and axis.has_title:
                info["axis_title_keys"] += placeholders_in(axis.axis_title.text_frame.text)
        except Exception:
            pass
    try:
        info["series"] = sum(1 for _ in chart.series)
    except Exception:
        pass
    return info


def walk_shapes(shapes):
    """Yield every shape, recursing into group shapes."""
    for sh in shapes:
        yield sh
        if getattr(sh, "shape_type", None) is not None and str(sh.shape_type) == "GROUP (6)":
            try:
                yield from walk_shapes(sh.shapes)
            except Exception:
                pass
        else:
            # Fallback: some group shapes expose .shapes without the GROUP type str.
            sub = getattr(sh, "shapes", None)
            if sub is not None and sub is not shapes:
                try:
                    yield from walk_shapes(sub)
                except Exception:
                    pass


# ----------------------------------------------------------------------------
# Probe one slide
# ----------------------------------------------------------------------------
def probe_slide(slide):
    """Return a dict describing one slide's keys and shapes."""
    placeholder_keys = []   # from {key}: text, image/video, chart/axis titles
    alttext_keys = []       # from alt text: tables and charts
    tables = {}             # alt text -> {"rows": n, "cols": n}
    shapes_out = []

    for sh in walk_shapes(slide.shapes):
        entry = {}
        try:
            entry["type"] = str(getattr(sh, "shape_type", "")) or ""
        except Exception:
            entry["type"] = ""
        entry["name"] = getattr(sh, "name", "") or ""

        # {key} placeholders in the shape's own text, or in its name (an image or
        # video shape has no text of its own, so its key often lives there)
        keys = []
        if getattr(sh, "has_text_frame", False):
            keys = placeholders_in(sh.text_frame.text)
        keys += [k for k in placeholders_in(entry["name"]) if k not in keys]
        entry["placeholder_keys"] = keys
        placeholder_keys += keys

        # alt text (binds tables and charts)
        title, descr = get_alt_text(sh)
        entry["alt_text"] = descr

        is_chart = bool(getattr(sh, "has_chart", False))
        is_table = bool(getattr(sh, "has_table", False))

        if is_chart:
            ci = chart_info(sh)
            entry["chart"] = ci
            entry["placeholder_keys"] += ci["title_keys"] + ci["axis_title_keys"]
            placeholder_keys += ci["title_keys"] + ci["axis_title_keys"]
            if descr:
                alttext_keys.append(descr)
        elif is_table:
            try:
                tbl = sh.table
                entry["table"] = {"rows": len(tbl.rows), "cols": len(tbl.columns)}
            except Exception:
                entry["table"] = {"rows": 0, "cols": 0}
            if descr:
                alttext_keys.append(descr)
                tables[descr] = entry["table"]

        shapes_out.append(entry)

    return {
        "placeholder_keys": sorted(set(placeholder_keys)),
        "alttext_keys": sorted(set(alttext_keys)),
        "tables": tables,
        "shapes": shapes_out,
    }


def probe(path):
    try:
        from pptx import Presentation
    except ImportError:
        die("python-pptx is not installed.\n"
            "  Run it through the launcher, which supplies it:\n"
            "    python scripts/run.py probe_template.py <template.pptx>")
    try:
        prs = Presentation(path)
    except Exception as e:
        die(f"could not open {path}: {e}")

    slides = [probe_slide(s) for s in prs.slides]
    return slides


# ----------------------------------------------------------------------------
# Rendering
# ----------------------------------------------------------------------------
def render_human(slides, numbers=None):
    """Print slides. `numbers` limits output to those 1-based slide numbers."""
    for i, s in enumerate(slides, 1):
        if numbers is not None and i not in numbers:
            continue
        print(paint(f'── Slide {i}  ("template": {i}) ', C.BOLD + C.CYAN)
              + paint("─" * 20, C.DIM))
        ph = s["placeholder_keys"]
        alt = s["alttext_keys"]
        print(f"  {paint('{key} placeholders', C.GREEN)}: "
              + (", ".join("{%s}" % k for k in ph) if ph else paint("(none)", C.DIM)))
        print(f"  {paint('alt-text keys (table/chart)', C.GREEN)}: "
              + (", ".join(alt) if alt else paint("(none)", C.DIM)))
        for sh in s["shapes"]:
            if "chart" in sh:
                ci = sh["chart"]
                bits = [f"chart type={ci['chart_type'] or '?'}", f"series={ci['series']}"]
                if sh["alt_text"]:
                    bits.append(f"alt='{sh['alt_text']}'")
                if ci["title_keys"]:
                    bits.append("title={%s}" % ",".join(ci["title_keys"]))
                if ci["axis_title_keys"]:
                    bits.append("axisTitle={%s}" % ",".join(ci["axis_title_keys"]))
                print(f"    {paint('▪', C.DIM)} " + "  ".join(bits))
            elif "table" in sh:
                t = sh["table"]
                alt = f" alt='{sh['alt_text']}'" if sh["alt_text"] else paint(" alt=(MISSING)", C.YELLOW)
                print(f"    {paint('▪', C.DIM)} table {t['rows']} rows x {t['cols']} cols{alt}")
        print()


def used_templates(data_path):
    """Template numbers a data.json refers to; None if it cannot be read (coverage reports why)."""
    try:
        with open(data_path, encoding="utf-8-sig") as f:
            data = json.load(f)
    except (OSError, ValueError):
        return None
    items = data.get("slides") if isinstance(data, dict) else data
    if not isinstance(items, list):
        return None
    return {d["template"] for d in items if isinstance(d, dict) and isinstance(d.get("template"), int)}


def coverage(slides, data_path):
    try:
        # utf-8-sig: Windows PowerShell 5.1 writes UTF-8 with a BOM.
        with open(data_path, encoding="utf-8-sig") as f:
            data = json.load(f)
    except FileNotFoundError:
        die(f"data file not found: {data_path}")
    except UnicodeDecodeError:
        die(f"{data_path} is not UTF-8 encoded. Save it as UTF-8 and run again.")
    except json.JSONDecodeError as e:
        die(f"{data_path} is not valid JSON: line {e.lineno}, col {e.colno}: {e.msg}", 2)

    data_slides = data.get("slides") if isinstance(data, dict) else data
    if not isinstance(data_slides, list):
        die("data.json has no `slides` array.")

    print(paint("══ Coverage (data.json vs template) ", C.BOLD) + paint("═" * 24, C.DIM))
    rc = 0
    for di, dslide in enumerate(data_slides, 1):
        if not isinstance(dslide, dict):
            continue
        t = dslide.get("template")
        label = f"data slide {di} → template {t}"
        if not isinstance(t, int) or not (1 <= t <= len(slides)):
            print(f"  {paint('✗', C.RED)} {label}: template number out of range "
                  f"(template has {len(slides)} slides)")
            rc = 1
            continue

        tmpl = slides[t - 1]
        tmpl_ph = set(tmpl["placeholder_keys"])
        tmpl_alt = set(tmpl["alttext_keys"])
        tmpl_all = tmpl_ph | tmpl_alt
        data_keys = {k for k in dslide if k != "template"}

        # Unfilled {key} placeholders LEAK literal text — this is the real failure.
        unfilled_ph = sorted(tmpl_ph - data_keys)
        # Unfilled chart/table just keep the template default — informational.
        unfilled_alt = sorted(tmpl_alt - data_keys)
        # Keys in data the template has nowhere to put — silently ignored.
        extra = sorted(data_keys - tmpl_all)

        # A table whose row/column count differs from the template's.
        size_notes = []
        for key, tdims in tmpl.get("tables", {}).items():
            value = dslide.get(key)
            if not (isinstance(value, dict) and isinstance(value.get("rows"), list)):
                continue
            rows = value["rows"]
            widths = sorted({len(r) for r in rows if isinstance(r, list)})
            if not widths:
                continue
            cols = str(widths[0]) if len(widths) == 1 else "%d-%d (rows differ)" % (widths[0], widths[-1])
            if len(rows) != tdims["rows"] or widths != [tdims["cols"]]:
                size_notes.append(f"`{key}`: template {tdims['rows']} rows x {tdims['cols']} cols, "
                                  f"data {len(rows)} rows x {cols} cols")

        if not unfilled_ph and not unfilled_alt and not extra and not size_notes:
            print(f"  {paint('✓', C.GREEN)} {label}: all keys matched")
            continue

        print(f"  {paint('•', C.CYAN)} {label}")
        for note in size_notes:
            print(f"      {paint('! table size differs from template', C.YELLOW)}: {note}")
        if unfilled_ph:
            rc = 1
            print(f"      {paint('✗ unfilled {key} (will leak as literal text)', C.RED)}: "
                  + ", ".join("{%s}" % k for k in unfilled_ph))
        if unfilled_alt:
            print(f"      {paint('! chart/table not provided (keeps template default)', C.YELLOW)}: "
                  + ", ".join(unfilled_alt))
        if extra:
            print(f"      {paint('! data key not in template (ignored)', C.YELLOW)}: "
                  + ", ".join(extra))
    print()
    return rc


# ----------------------------------------------------------------------------
# Main
# ----------------------------------------------------------------------------
def main():
    p = argparse.ArgumentParser(description="Probe a SlidePack template.pptx for the keys it expects.")
    p.add_argument("pptx", nargs="?", default="template.pptx", help="template PPTX (default: template.pptx)")
    p.add_argument("--data", help="data.json to coverage-check against the template")
    p.add_argument("--slide", metavar="N[,N...]",
                   help="only show these 1-based slide numbers (the value of `template` in data.json)")
    p.add_argument("--all", action="store_true",
                   help="with --data, list every slide instead of only the ones the data uses")
    p.add_argument("--json", action="store_true", help="emit machine-readable JSON")
    args = p.parse_args()

    if not os.path.isfile(args.pptx):
        die(f"file not found: {args.pptx}")

    slides = probe(args.pptx)

    numbers = None
    if args.slide:
        try:
            numbers = {int(x) for x in args.slide.split(",")}
        except ValueError:
            die(f"--slide expects numbers such as 15 or 3,15, got {args.slide!r}", 2)
        bad = sorted(n for n in numbers if not 1 <= n <= len(slides))
        if bad:
            die(f"--slide {bad}: the template has {len(slides)} slides (numbered from 1)", 2)
    elif args.data and not args.all:
        numbers = used_templates(args.data)

    if args.json:
        shown = [dict(s, number=i) for i, s in enumerate(slides, 1) if numbers is None or i in numbers]
        print(json.dumps({"slides": shown}, ensure_ascii=False, indent=2))
        if args.data:
            print("note: --json lists the template's keys only; run without --json "
                  "for the coverage check against --data.", file=sys.stderr)
        sys.exit(0)

    render_human(slides, numbers)
    if args.data:
        sys.exit(coverage(slides, args.data))
    sys.exit(0)


if __name__ == "__main__":
    main()
